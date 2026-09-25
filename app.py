import io
from flask import Flask, render_template, request, send_file
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

app = Flask(__name__)

def formato_consistencia_conta_teso_grupo(nombre_archivo):
    try :
            df = pd.read_excel(nombre_archivo, 
                            skiprows = 16, 
                            skipfooter = 1, 
                            header = None,
                            dtype = object,
                            engine = 'xlrd')
    except:
            nombre_archivo.seek(0)
            df = pd.read_excel(nombre_archivo, 
                            skiprows = 16, 
                            skipfooter = 1, 
                            header = None,
                            dtype = object,
                            engine = 'xlrd')
    df = df.dropna(axis = 1, how = 'all')
    df =  df.drop(df.columns[-1], axis = 1)
    lista_cols = ['tipo', 'nro_asiento', 'entidad_conta', 'da_conta', 'documento_conta', 'tipo_conta', 't_conta','r_conta','monto_conta',
        'entidad_teso', 'da_teso', 'documento_teso','t_teso', 'r_teso', 'codigo_op', 'monto_teso']
    cols = {0:'tipo', 3:'nro_asiento', 9:'entidad_conta', 11:'da_conta', 14:'documento_conta', 20:'tipo_conta', 28:'t_conta',30:'r_conta',36:'monto_conta',
        45:'entidad_teso', 46:'da_teso', 49:'documento_teso', 54:'t_teso', 59:'r_teso', 60:'codigo_op', 62:'monto_teso'}
    df = df.rename(columns = cols)
    df = df.reindex(columns = lista_cols)
    conta = df.loc[:,:'monto_conta'].copy()
    teso = df.loc[:,'entidad_teso':'monto_teso'].copy()
    conta['key'] = 'Entidad:'+ conta['entidad_conta'].astype(str) + 'DA:' + conta['da_conta'].astype(str) + 'Comprobante:' + conta['documento_conta']
    teso['key'] = 'Entidad:'+ teso['entidad_teso'].astype(str) + 'DA:' +teso['da_teso'].astype(str) + 'Comprobante:' + teso['documento_teso']
    contraste = conta.merge(right = teso, how = 'outer', on = 'key')
    contraste['monto_conta'] = contraste['monto_conta'].fillna(0)
    contraste['monto_teso'] = contraste['monto_teso'].fillna(0)
    contraste['diferencia'] = contraste['monto_conta'] - contraste['monto_teso']
    contraste = contraste[contraste['key'].notna()]
    output = io.BytesIO()
    return archivo_final(output, contraste)

def analisis_contabilidad(nombre_archivo):
        pass

def reporte_ejecucion_recursos_estructura(nombre_archivo):
    df = pd.read_excel(nombre_archivo, 
                       engine = 'xlrd')
    filas_encontradas = pd.Series(False, index=df.index)
    for col in df.columns:
        coincidencias = df[col].astype(str).str.contains(pat = '^Rubro:$', na = False, case = False)
        filas_encontradas = filas_encontradas | coincidencias
    tabla_rubros = pd.DataFrame({'rubros':df.loc[filas_encontradas,df.columns[49]]})
    tabla_ent_transferencia = pd.DataFrame({'ent_transf': df.loc[df[df.columns[50]].notna(), df.columns[50]]}).set_index(tabla_rubros.index)
    indices = list(tabla_rubros.index)
    indices.append(len(df))
    rubros = list(tabla_rubros['rubros'])
    entidad_transf = list(tabla_ent_transferencia['ent_transf'])
    rangos = ['rango' + str(i) for i in range(1, len(list(tabla_rubros['rubros'])) + 1)]
    rubro_rangos = dict(zip(rangos, rubros))
    ent_transf_rangos = dict(zip(rangos, entidad_transf))
    df['rangos'] = pd.cut(df.index, bins = indices, right = True, labels = rangos)
    df_tratado = pd.DataFrame()
    for rango in rangos:
        df_temp = df[df['rangos'] == rango].iloc[3:-4,:].dropna(axis = 1, how = 'all')
        df_temp = df_temp[df_temp.loc[:,df_temp.columns[-2]].notna()].dropna(axis = 1, how = 'all')
        df_tratado = pd.concat(objs = [df_tratado, df_temp], axis = 0, ignore_index = False)
    df_tratado['rubro'] = df_tratado['rangos'].map(rubro_rangos)
    df_tratado['ent_transferncia'] = df_tratado['rangos'].map(ent_transf_rangos)
    df_tratado = df_tratado.drop(columns = 'rangos')
    df_tratado.columns = ['doc_dev', 'doc_perc', 'sec', 'tipo_doc', 'fecha_aprobacion', 'devengado', 'percibido', 'resumen', 'rubro', 'ent_transf']
    df_tratado['fecha_aprobacion'] = pd.to_datetime(df_tratado['fecha_aprobacion']).dt.strftime('%d/%m/%y')
    output = io.BytesIO()
    return archivo_final(output, df_tratado)

def consulta_partidas(nombre_archivo):
    df = pd.read_excel(nombre_archivo)
    objeto = df.iloc[1,21]
    ent_transferencia = df.iloc[2, 21]
    df = df.iloc[6:,:].dropna(axis = 1, how = 'all')
    df = df[df[df.columns[9]].notna()].dropna(axis = 1, how = 'all')
    df.columns = ['clase_gasto', 'preventivo', 'compromiso', 'devengado', 'pago', 'sec', 'fecha_elaboracion',
                  'fecha_verificacion', 'fecha_aprobacion', 'glosa', 'importe', 'multas', 'total_autorizado',
                  'retenciones', 'liquido']
    df = df.drop(['fecha_elaboracion', 'fecha_verificacion'], axis = 1)
    df['fecha_aprobacion'] = pd.to_datetime(df['fecha_aprobacion']).dt.strftime('%d/%m/%Y')
    df['objeto'] = objeto
    df['ent_transferencia'] = ent_transferencia
    output = io.BytesIO()
    return archivo_final(output, df)

def archivo_final(salida, df):
    with pd.ExcelWriter(salida, engine = 'openpyxl') as writer:
            df.to_excel(writer, index = False, sheet_name = 'Procesados')
            worksheet = writer.sheets['Procesados']
            header_font = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
            header_fill = PatternFill(start_color='1F4E78', end_color='1F4E78', fill_type='solid') # Azul oscuro            
            center_alignment = Alignment(horizontal='center', vertical='center')
            for col_num in range(1, len(df.columns) + 1):
                    cell = worksheet.cell(row=1, column=col_num)
                    cell.font = header_font
                    cell.fill = header_fill
                    cell.alignment = center_alignment
            for col in worksheet.columns:
                    max_len = max(len(str(cell.value or '')) for cell in col)
                    col_letter = get_column_letter(col[0].column)
                    worksheet.column_dimensions[col_letter].width = max(max_len + 4, 12)
    salida.seek(0)
    return salida

tipos_reportes = {
        'analisis_consistencia' : formato_consistencia_conta_teso_grupo,
        'reporte_rec_estructura' : reporte_ejecucion_recursos_estructura,
        'reporte_consulta_partidas' : consulta_partidas
}

@app.route('/')
def index():
    return render_template('index.html')
@app.route('/transform', methods = ['POST'])
def transform():
    report_type = request.form.get('report_type')
    file = request.files.get('excel_file')
    if file.filename == '':
            return 'Nombre de archivo no valido', 400
    if report_type not in tipos_reportes:
            return "Tipo de reporte no válido.", 400
    funcion_procesadora = tipos_reportes[report_type]
    processed_file = funcion_procesadora(file)
    output_filename = f"procesado_{file.filename.rsplit('.', 1)[0]}.xlsx"
    return send_file(
                    processed_file,
                    mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                    as_attachment=True,
                    download_name=output_filename
                )
if __name__ == '__main__':
    app.run(debug = True)